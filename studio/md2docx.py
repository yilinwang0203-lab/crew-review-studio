#!/usr/bin/env python3
"""通用 Markdown → Word(.docx) 转换器（纯本地离线，零联网）。

支持：# 标题(1-3级)、表格(| a | b |)、- 无序列表、> 引用、`代码`行内、
普通段落、空行。够用即可，不追求完整 Markdown 规范。

用法：python3 md2docx.py 输入.md [输出.docx]
不给输出名则与输入同名换 .docx。
"""
import sys, re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

# 中日文字体：必须被 fontconfig 索引(soffice 才渲染得出)且含中文+日文假名。
# 实测 PingFang SC 不在 fc-list→soffice 回退方框；Heiti SC(华文黑体 STHeiti)
# 在 fc-list 且含中文与假名(の/中)，Word for Mac 也有，故用它。
CJK_FONT = 'Heiti SC'


def _set_run_cjk(run, name=CJK_FONT):
    """设 run 的中西文字体，关键是 w:eastAsia（python-docx 的 font.name 不设它，中文会方框）。"""
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), name)


def _apply_cjk_font(doc, name=CJK_FONT):
    """全文档统一中文字体：样式 + 所有段落/表格单元格的 run。"""
    for sname in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3',
                  'List Bullet', 'List Number', 'No Spacing', 'Title'):
        try:
            st = doc.styles[sname]
            st.font.name = name
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find(qn('w:rFonts'))
            if rf is None:
                rf = rpr.makeelement(qn('w:rFonts'), {})
                rpr.insert(0, rf)
            for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                rf.set(qn(attr), name)
        except KeyError:
            pass

    def _walk(paras):
        for p in paras:
            for r in p.runs:
                _set_run_cjk(r, name)
    _walk(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                _walk(cell.paragraphs)


def add_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i in range(len(header)):
            cells[i].text = r[i] if i < len(r) else ''
    doc.add_paragraph('')


def split_row(line):
    # | a | b | c |  ->  [a, b, c]
    parts = line.strip().strip('|').split('|')
    return [p.strip() for p in parts]


def clean_inline(s):
    # 去掉行内 markdown 强调标记，保留文字
    s = re.sub(r'`([^`]*)`', r'\1', s)
    s = re.sub(r'\*\*([^*]*)\*\*', r'\1', s)
    return s


def convert(md_path, out_path):
    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    doc = Document()
    doc.styles['Normal'].font.size = Pt(10.5)

    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            doc.add_paragraph(line, style='No Spacing')
            i += 1
            continue

        # 表格：当前行是 | ... | 且下一行是分隔行 |---|
        if stripped.startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1]):
            header = [clean_inline(c) for c in split_row(stripped)]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([clean_inline(c) for c in split_row(lines[i].strip())])
                i += 1
            add_table(doc, header, rows)
            continue

        if not stripped:
            i += 1
            continue

        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(clean_inline(m.group(2)), level=level)
        elif stripped.startswith('> '):
            p = doc.add_paragraph(clean_inline(stripped[2:]))
            p.runs[0].italic = True if p.runs else None
        elif re.match(r'^[-*]\s+', stripped):
            doc.add_paragraph(clean_inline(re.sub(r'^[-*]\s+', '', stripped)), style='List Bullet')
        elif re.match(r'^\d+\.\s+', stripped):
            doc.add_paragraph(clean_inline(re.sub(r'^\d+\.\s+', '', stripped)), style='List Number')
        else:
            doc.add_paragraph(clean_inline(stripped))
        i += 1

    _apply_cjk_font(doc)
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else str(Path(src).with_suffix('.docx'))
    out = convert(src, dst)
    print(f"已生成: {out}")
